"""Application XLSX export: frozen values, provenance and formula-safe text."""
import io
from collections import defaultdict
from datetime import datetime
from zoneinfo import ZoneInfo
from .core import metrics
from .checks import upload_summary


def monthly_xlsx(month,rows,batches,period=None,history=()):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    period=period or dict(status='review',revision=0)
    all_rows=rows
    rows=[r for r in rows if r['month']==month]
    result=metrics(rows)
    book=Workbook();book.remove(book.active)
    def sheet(name,headers,values,widths=None):
        ws=book.create_sheet(name);ws.append(headers)
        for cell in ws[1]:
            cell.fill=PatternFill('solid',fgColor='16324F');cell.font=Font(color='FFFFFF',bold=True)
        for record in values:
            ws.append(record)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value,str):
                    cell.data_type='s'  # User merchant names never become Excel formulas.
                elif isinstance(cell.value,(int,float)) and not isinstance(cell.value,bool):
                    cell.number_format='#,##0;[Red](#,##0);"–"'
                cell.alignment=Alignment(vertical='top',wrap_text=True)
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;ws.sheet_view.showGridLines=False
        for n,header in enumerate(headers,1):
            ws.column_dimensions[get_column_letter(n)].width=(widths or {}).get(header,22)
        ws.row_dimensions[1].height=28
        return ws
    status='확정' if period['status']=='finalized' else '검토 중'
    sheet('월정산 요약',['항목','값'],[
      ['정산월',month],['상태',status],['확정 차수',period.get('revision',0)],['확정일시',period.get('finalized_at','')],
      ['월정산 수입',result['income']],['사업 지출 (카드 환불 차감)',result['expense']],['최종 순수익 (월정산 기준)',result['residual']],
      ['은행 순현금 증감',result['bank_net']],['확인 필요 거래',result['pending']],['전체 거래 수',len(rows)],
      ['작성일시',datetime.now(ZoneInfo('Asia/Seoul')).isoformat()],
      ['계산 기준','사업 은행 입금 − 사업 은행 지출 − 카드 사용액(환불 차감). 카드대금·이체·대출·개인/제외는 사업 집계 제외.'],
      ['안내','최종 순수익은 월정산 잔여금이며 재고·미수금·감가상각 등을 반영한 회계상 순이익과 다릅니다. 검토 중 자료는 확정 보고서가 아닙니다.']],{'항목':35,'값':95})
    names={b['id']:b.get('filename','') for b in batches}
    headers=['거래일','정산월','자료','계좌/카드','거래처/내역','카테고리','처리 구분','입금/환불','출금/사용','포함','확인','반영 기준','납부명세서 월','할부 회차','총 회차','원본 파일','원본 시트','원본 행','분류 근거']
    sheet('거래 상세',headers,[[r['date'],r['month'],r['source'],r['account'],r['description'],r['category'],r['treatment'],
         r['inflow'],r['outflow'],r['include'],r['reviewed'],r.get('expense_basis',r.get('settlement_basis','거래일')),
         r.get('billing_month',''),r.get('installment_number'),r.get('installment_total'),names.get(r['batch_id'],''),r.get('source_sheet',''),r.get('row_number'),r.get('reason','')] for r in rows],{'거래처/내역':42,'원본 파일':55,'분류 근거':55})
    grouped=defaultdict(lambda:[0,0,0])
    for r in rows:
        if r['include'] and r['treatment']=='사업':
            grouped[r['category']][0]+=r['inflow'] if r['source']=='은행' else 0
            grouped[r['category']][1]+=r['outflow']-(r['inflow'] if r['source']=='카드' else 0)
            grouped[r['category']][2]+=1
    sheet('카테고리 집계',['카테고리','사업 수입','사업 지출 (환불 차감)','거래 수'],[[cat,*v] for cat,v in sorted(grouped.items())])
    relevant={r['batch_id'] for r in rows}
    summaries=upload_summary([b for b in batches if b['id'] in relevant],all_rows)
    for summary,batch in zip(summaries,[b for b in batches if b['id'] in relevant]):
        summary['본 보고서월 거래 수']=sum(r['batch_id']==batch['id'] for r in rows)
    if summaries:
        sheet('업로드 확인표',list(summaries[0]),[list(s.values()) for s in summaries],{'파일':55,'구매/거래 기간':32})
    events=[]
    for e in history:
        if e.get('action')=='edit':
            previous={r['id']:r for r in e.get('before',[])}
            for r in e.get('after',[]):
                if r.get('month')!=month:continue
                before=previous.get(r['id'],{})
                for field in ('category','treatment','include','reviewed','card_account'):
                    if before.get(field)!=r.get(field):
                        events.append([e['created_at'],'내역 수정',r['description'],field,str(before.get(field,'')),str(r.get(field,'')),''])
        elif e.get('month')==month:
            events.append([e['created_at'],{'finalize':'정산 확정','reopen':'확정 해제'}.get(e['action'],e['action']),'','','','',e.get('reason','')])
    sheet('수정 이력',['일시','작업','거래 내역','변경 항목','이전 값','변경 값','사유'],events,{'거래 내역':42,'사유':55,'일시':35})
    data=io.BytesIO();book.save(data);return data.getvalue()
