import io
import json
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from settlement.store import Store,ConflictError,restore_to_new_sqlite
from settlement.workflow import ClosedPeriodError,row_digest
from settlement.report import monthly_xlsx
from test_settlement import sample,confirmed

class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.url='sqlite:///'+str(Path(self.tmp.name)/'workflow.sqlite3');self.store=Store(self.url)
        self.row=sample();self.context=dict(kind='upload',month=self.row['month'],batches=[{'id':self.row['batch_id']}])
    def save_row(self):
        row=dict(self.row,reviewed=True,category='사입')
        self.store.import_rows(self.context['batches'],[row])
        return self.store.rows()[0]
    def test_pending_draft_survives_connection_and_version_noop(self):
        first=self.store.save_draft('d',[self.row],self.context)
        other=Store(self.url);self.assertEqual(other.draft('d')['rows'][0]['category'],'미분류')
        self.assertEqual(other.save_draft('d',[self.row],self.context,first['version'])['version'],1)
        self.assertEqual(len(other.draft_versions('d')),1)
        self.assertTrue(other.backups())
    def test_conflicting_autosave_is_not_last_writer_wins(self):
        self.store.save_draft('d',[self.row],self.context)
        updated=dict(self.row,category='사입')
        self.store.save_draft('d',[updated],self.context,1)
        with self.assertRaises(ConflictError):self.store.save_draft('d',[dict(updated,category='광고비')],self.context,1)
        self.assertEqual(self.store.draft('d')['rows'][0]['category'],'사입')
    def test_previous_draft_restores_as_new_version(self):
        self.store.save_draft('d',[self.row],self.context)
        self.store.save_draft('d',[dict(self.row,category='사입')],self.context,1)
        restored=self.store.restore_draft_version('d',1,2)
        self.assertEqual(restored['version'],3);self.assertEqual(restored['rows'][0]['category'],'미분류')
        self.assertEqual(len(self.store.draft_versions('d')),3)
    def test_import_consumes_draft_atomically(self):
        self.store.save_draft('d',[self.row],self.context)
        row=dict(self.row,reviewed=True,category='사입')
        with self.assertRaises(ConflictError):self.store.import_rows(self.context['batches'],[row],draft_id='d',draft_version=0)
        self.assertFalse(self.store.rows());self.assertEqual(self.store.draft('d')['status'],'active')
        self.store.import_rows(self.context['batches'],[row],draft_id='d',draft_version=1)
        self.assertEqual(self.store.draft('d')['status'],'saved');self.assertFalse(self.store.drafts())
    def test_finalize_locks_edit_import_and_draft(self):
        row=self.save_row();period=self.store.finalize_month(row['month'],row_digest([row]),upload_confirmed=True)
        self.assertEqual(period['status'],'finalized')
        with self.assertRaises(ClosedPeriodError):self.store.edit_rows([dict(row,category='세금')])
        with self.assertRaises(ClosedPeriodError):self.store.save_draft('d',[row],self.context)
        new=confirmed(file='new')
        with self.assertRaises(ClosedPeriodError):self.store.import_rows([{'id':new['batch_id']}],[new])
        self.assertEqual(self.store.rows()[0]['category'],'사입')
    def test_finalization_blocks_unfinished_jobs_and_stale_rows(self):
        row=self.save_row()
        with self.assertRaises(ConflictError):self.store.finalize_month(row['month'],'old',upload_confirmed=True)
        self.store.save_draft('d',[row],self.context)
        with self.assertRaisesRegex(ValueError,'임시 작업'):self.store.finalize_month(row['month'],row_digest([row]),upload_confirmed=True)
        self.store.archive_draft('d',1,'테스트 작업 종료')
        self.store.finalize_month(row['month'],row_digest([row]),upload_confirmed=True)
    def test_reopen_requires_reason_and_revision_and_keeps_history(self):
        row=self.save_row();p=self.store.finalize_month(row['month'],row_digest([row]),upload_confirmed=True)
        with self.assertRaises(ValueError):self.store.reopen_month(row['month'],'',p['version'])
        self.store.reopen_month(row['month'],'택배 분류 정정',p['version'])
        self.store.edit_rows([dict(row,category='택배비')]);rows=self.store.rows()
        second=self.store.finalize_month(row['month'],row_digest(rows),upload_confirmed=True)
        self.assertEqual(second['revision'],2)
        events=self.store.month_history(row['month']);self.assertEqual(sum(e['action']=='finalize' for e in events),2)
        self.assertTrue(any(e.get('reason')=='택배 분류 정정' for e in events))
    def test_restore_backup_preserves_pending_work_and_closed_month(self):
        row=self.save_row();self.store.finalize_month(row['month'],row_digest([row]),upload_confirmed=True)
        other=sample(file='october');other.update(month='2026-10',date='2026-10-01')
        self.store.save_draft('d',[other],dict(kind='upload',batches=[{'id':other['batch_id']}]))
        restored=restore_to_new_sqlite(self.store.export(),Path(self.tmp.name)/'restored.sqlite3')
        self.assertEqual(restored.period(row['month'])['status'],'finalized')
        self.assertEqual(restored.draft('d')['rows'][0]['category'],'미분류')
        self.assertEqual(restored.draft_versions('d'),self.store.draft_versions('d'))
        self.assertEqual(restored.backups(),self.store.backups())
    def test_recovery_point_has_verified_payload_and_can_restore_separately(self):
        row=self.save_row();ident=self.store.checkpoint('테스트 복구본')
        restored=restore_to_new_sqlite(self.store.backup_bytes(ident),Path(self.tmp.name)/'point.sqlite3')
        self.assertEqual(restored.rows(),self.store.rows())
    def test_source_file_counts_must_match_before_finalization(self):
        row=dict(self.row,reviewed=True,category='사입')
        self.store.import_rows([dict(id=row['batch_id'],source_rows=2,error_count=0)],[row])
        with self.assertRaisesRegex(ValueError,'원본 행 수'):self.store.finalize_month(row['month'],row_digest(self.store.rows()),upload_confirmed=True)
    def test_migration_only_allows_empty_target_and_preserves_pending_state(self):
        from tools.migrate_empty import copy_to_empty_store
        self.store.save_draft('d',[self.row],self.context)
        target=Store('sqlite:///'+str(Path(self.tmp.name)/'empty.sqlite3'))
        data=json.loads(self.store.export())
        copy_to_empty_store(target,data,'test-digest')
        self.assertEqual(target.draft('d')['rows'][0]['category'],'미분류')
        with self.assertRaisesRegex(ValueError,'비어 있지'):copy_to_empty_store(target,data,'test-digest')
        self.assertEqual(len(target.drafts()),1)

    def test_excel_snapshot_totals_and_formula_safe_merchant(self):
        row=dict(self.row,reviewed=True,category='사입',description='=HYPERLINK("https://example.invalid")')
        payload=monthly_xlsx(row['month'],[row],[{'id':row['batch_id'],'filename':'test.xlsx'}])
        book=load_workbook(io.BytesIO(payload),data_only=False)
        summary=dict(book['월정산 요약'].iter_rows(min_row=2,values_only=True))
        self.assertEqual(summary['사업 지출 (카드 환불 차감)'],row['outflow'])
        self.assertEqual(summary['최종 순수익 (월정산 기준)'],-row['outflow'])
        self.assertEqual(book['거래 상세']['E2'].data_type,'s')
        self.assertEqual(book['거래 상세']['E2'].value,row['description'])

if __name__=='__main__':unittest.main()
